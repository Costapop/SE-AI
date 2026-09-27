#!/usr/bin/env python3
"""Записывает examples/examples.json в девять файлов .xlsx (лист «Спектр» + лист «Условия») и examples/answer_key.xlsx с истиной.
Требует openpyxl. Запуск: python3 tools/write_xlsx.py"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
ex = json.load(open(os.path.join(ROOT, "examples/examples.json"), encoding="utf-8"))
MAT = {"tio2": "TiO2 (аморфный)", "ta2o5": "Ta2O5 (аморфный)", "nb2o5": "Nb2O5 (аморфный)", "si3n4": "Si3N4 (LPCVD)", "hfo2": "HfO2 (аморфный)",
       "sinx": "SiNx:H (PECVD, обогащённый Si; край поглощения в диапазоне)", "ceo2": "CeO2 (поликристаллический; край поглощения в диапазоне)", "wo3": "WO3 (тонкая плёнка; край поглощения в диапазоне)",
       "zro2": "ZrO2 (аморфный)", "sio2": "SiO2 (термический)", "al2o3": "Al2O3 (аморфный)"}
SUB = {"bk7": "Стекло BK7 (Зельмейер Schott)", "silica": "Плавленый кварц (Malitson 1965)", "si": "Кремний c-Si (Green 2008, 250–1450 нм)", "si_aspnes": "Кремний c-Si (Aspnes–Studna 1983)"}
REF = {"tio2": dict(A=255.83, E0=4.00, C=1.77, Eg=3.40, Auv=137.65, Euv=11.0), "ta2o5": dict(A=321.34, E0=5.30, C=2.60, Eg=4.20, Auv=11.33, Euv=12.0),
       "nb2o5": dict(A=307.24, E0=4.70, C=2.30, Eg=3.75, Auv=69.07, Euv=11.0), "si3n4": dict(A=156.51, E0=7.20, C=3.60, Eg=4.60, Auv=103.31, Euv=13.0),
       "hfo2": dict(A=396.97, E0=6.20, C=2.80, Eg=5.20, Auv=16.82, Euv=13.0),
       "sinx": dict(A=52.04, E0=7.38, C=3.26, Eg=2.17, Auv=79.35, Euv=13.0), "ceo2": dict(A=65.58, E0=3.92, C=0.89, Eg=2.81, Auv=414.58, Euv=13.0),
       "wo3": dict(A=57.22, E0=4.41, C=1.20, Eg=3.26, Auv=311.01, Euv=13.0),
       "zro2": dict(A=449.00, E0=5.95, C=3.00, Eg=5.00, Auv=28.74, Euv=13.0), "sio2": dict(A=295.50, E0=10.35, C=3.00, Eg=8.70, Auv=10.76, Euv=16.0), "al2o3": dict(A=182.48, E0=9.50, C=3.50, Eg=6.60, Auv=11.45, Euv=14.0)}
bold = Font(name="Arial", bold=True); reg = Font(name="Arial"); hdr_fill = PatternFill("solid", fgColor="DDE7F0")
for e in ex:
    wb = Workbook(); ws = wb.active; ws.title = "Спектр"
    ws.append(["λ, нм", "Ψ, град", "Δ, град", "σΨ, град", "σΔ, град"])
    for c in range(1, 6):
        ws.cell(1, c).font = bold; ws.cell(1, c).fill = hdr_fill; ws.cell(1, c).alignment = Alignment(horizontal="center")
    for i in range(len(e["lam"])):
        ws.append([e["lam"][i], e["psi"][i], e["del"][i], e["sigPsi"][i], e["sigDel"][i]])
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.font = reg; c.number_format = "0.0000"
    for c, w in zip("ABCDE", (10, 12, 12, 12, 12)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A2"
    w2 = wb.create_sheet("Условия")
    rows = [("Образец", e["title"]), ("Материал плёнки (априори)", MAT[e["mat"]]),
            ("Модель дисперсии", "Тауц–Лоренц + УФ-полюс, ε∞ = 1 (физически допустимая, КК-согласованная)"),
            ("Справочные параметры", "A = %.2f, E0 = %.2f эВ, C = %.2f эВ, Eg = %.2f эВ, A_uv = %.2f эВ², E_uv = %.1f эВ" % tuple(REF[e["mat"]][k] for k in ("A", "E0", "C", "Eg", "Auv", "Euv"))),
            ("Подложка", SUB[e["sub"]] + ("; плоскопараллельная, толщина %g мм" % (e["back"]["ds"] / 1e6) if e.get("back") else "; задней стороны нет (полубесконечная)")),
            ("Задняя сторона подложки", ("учитывается: доля отражённого задней гранью света, попадающего в детектор, %.2f ± %.2f (независимое измерение по голой подложке; в фите не варьируется)" % (e["back"]["f"], e["back"]["ferr"])) if e.get("back") else "нет (зачернена / матирована / отсечена)"),
            ("Угол падения, град", e["phi"]),
            ("Толщина (номинал), нм", e["dNom"]), ("Неопределённость толщины, %", "±%d" % e["dRange"]),
            ("Диапазон длин волн, нм", "%g – %g, шаг %g" % (e["lam"][0], e["lam"][-1], e["lam"][1] - e["lam"][0])),
            ("Спектральная полоса прибора, нм (FWHM)", e["bw"]), ("Шумовая модель", "столбцы σΨ, σΔ — повторяемость прибора (σΔ растёт при Ψ → 0)"),
            ("Прибор", "RCE-эллипсометр с ПЗС, все длины волн одновременно; Δ в диапазоне 0–360°; конвенция N = n − ik (Аззам–Башара)"),
            ("Возможные дефекты", "неоднородность по глубине (линейный градиент n), шероховатость поверхности (EMA-слой Бруггемана 50/50), поглощение (хвост ниже края: второй осциллятор Тауца–Лоренца с теми же E0, C и порогом E_t = 1.2 эВ; размер — k хвоста при 400 нм)"),
            ("Назначение", "тестовый файл для приложения: опорный фит однородной плёнки → ИИ по остатку → окончательный фит; истинные параметры — в answer_key.xlsx")]
    for k, v in rows:
        w2.append([k, v])
    for r in w2.iter_rows():
        r[0].font = bold; r[1].font = reg; r[1].alignment = Alignment(wrap_text=True, vertical="top")
    w2.column_dimensions["A"].width = 40; w2.column_dimensions["B"].width = 95
    wb.save(os.path.join(ROOT, "examples", e["file"]))
wb = Workbook(); ws = wb.active; ws.title = "Истина"
heads = ["Файл", "Образец", "d, нм", "Градиент δ, % (n растёт к поверхности > 0)", "EMA-слой шероховатости, нм", "Поглощение: k хвоста при 400 нм, ×10⁻³", "Задняя сторона: истинная доля (задано ± неопр.)", "Фактический угол, град", "Полоса, нм", "Сдвиг шкалы λ, нм",
         "A", "E0, эВ", "C, эВ", "Eg, эВ", "A_uv", "A_k хвоста, эВ", "n(450)", "n(550)", "n(650)", "n(750)", "k(450)", "k(550)"]
ws.append(heads)
for c in range(1, len(heads) + 1):
    ws.cell(1, c).font = bold; ws.cell(1, c).fill = hdr_fill; ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="top")
for e in ex:
    t = e["truth"]; m = t["mat"]; kk = t.get("k", {})
    bk = ("%.3f (%.2f ± %.2f)" % (t["back"]["f"], e["back"]["f"], e["back"]["ferr"])) if t.get("back") else "нет"
    ws.append([e["file"], e["title"], t["d"], 100 * t["delta"], t["dRough"], 1000 * t.get("k400", 0), bk, t["phiTrue"], t["bw"], t["lamOffset"], round(m["A"], 2), m["E0"], m["C"], m["Eg"], round(m["Auv"], 2), round(m.get("Ak", 0), 4)]
              + [round(t["n"][k], 4) for k in ("450", "550", "650", "750")] + [round(kk.get(k, 0), 5) for k in ("450", "550")])
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.font = reg
for c in range(1, len(heads) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 16
ws.column_dimensions["A"].width = 36; ws.column_dimensions["B"].width = 34
ws.append([]); ws.append(["Примечание", "Дисперсия каждой плёнки отличается от справочной (силы осцилляторов, E0, C, Eg), как у реального образца; угол отличается от номинального на 0.01–0.03°; дрейф и шум добавлены. "
                                        "Поглощение задаётся хвостом ниже края (второй осциллятор Тауца–Лоренца с теми же E0, C, порог E_t = 1.2 эВ, сила A_k); k(450), k(550) — полный k плёнки (край материала + хвост); у образцов 01–05 хвоста нет. "
                                        "Задняя сторона подложки (образец 09): пучки, отражённые задней гранью плоскопараллельной подложки, складываются с передним некогерентно; доля собранного света задана из независимого измерения и не фитируется."])
ws.cell(ws.max_row, 1).font = bold; ws.cell(ws.max_row, 2).font = reg
wb.save(os.path.join(ROOT, "examples", "answer_key.xlsx"))
print("xlsx записаны")
